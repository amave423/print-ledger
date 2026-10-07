"""Перенос данных из старой таблицы «отработка принтера.xlsx».

Создаёт:
  seed.sql        — вставить в Supabase SQL Editor после создания пользователя
  demo-data.json  — для локального демо-режима

Запуск: python tools/import_xlsx.py "путь/к/отработка принтера.xlsx" [папка_вывода]
Оба файла содержат ваши данные и не должны попадать в репозиторий (см. .gitignore).
"""
import json
import re
import sys
import uuid
from pathlib import Path

import openpyxl

PRINTER = {"name": "Qidi Plus 4", "price": 36400, "bought_on": "2025-12-01"}
CATEGORIES = [
    (r"сопл", "Сопла"), (r"сушилк", "Оборудование"), (r"резак|ключ|штанг", "Инструмент"),
    (r"пакет|силикагел", "Хранение"), (r"носок|носки|вентилятор|пластин", "Запчасти"),
]


def category(name):
    low = name.lower()
    return next((c for pat, c in CATEGORIES if re.search(pat, low)), "Расходники")


def iso(v):
    return v.date().isoformat()


def main():
    src = Path(sys.argv[1])
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(".")
    wb = openpyxl.load_workbook(src)
    printer_id = str(uuid.uuid4())

    spools = []
    for name_brand, weight, price, _, bought in wb["Пластик"].iter_rows(min_row=2, max_col=5, values_only=True):
        if not name_brand:
            continue
        name, _, brand = (x.strip() for x in name_brand.partition(","))
        spools.append({"id": str(uuid.uuid4()), "name": name, "brand": brand, "weight": weight, "price": price, "bought_on": iso(bought)})

    extras = []
    for bought, name, price, *_ in wb["Допы|Расходники"].iter_rows(min_row=2, max_col=4, values_only=True):
        if not name:
            continue
        extras.append({"id": str(uuid.uuid4()), "name": name.strip(), "category": category(name), "price": price, "bought_on": iso(bought)})

    # в старой таблице пластик указан типом, а конкретная катушка — ссылкой в формуле выручки
    spool_rows = {i + 2: s for i, s in enumerate(spools)}
    orders = []
    for row in wb["Печать"].iter_rows(min_row=2, max_col=7):
        date, plastic, revenue, weight, model, prnt = (c.value for c in row[:6])
        if not date:
            continue
        m = re.search(r"Пластик!D(\d+)", str(revenue))
        ref = spool_rows.get(int(m.group(1))) if m else None
        # в таблице тип пластика указан общим словом, катушка — ссылкой в формуле;
        # для старых заказов берём тип (у TPU — с твёрдостью), списание идёт по всем катушкам этого типа
        if ref:
            words = ref["name"].split()
            plastic = " ".join(words[:2]) if words[0].upper() == "TPU" else words[0]
        orders.append({"id": str(uuid.uuid4()), "ordered_on": iso(date), "plastic": plastic, "weight": weight,
                       "waste": 0, "print_price": prnt, "model_price": model or 0, "prepay": 0, "hours": 0,
                       "item": "", "client": "", "note": "", "printer_id": printer_id, "temps": None})

    printers = [{"id": printer_id, **PRINTER, "power": None}]
    data = {"printers": printers, "spools": spools, "extras": extras, "orders": orders,
            "settings": {"tariff": 6.5, "presets": {}}}
    (out / "demo-data.json").write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")

    def q(v):
        if v is None:
            return "null"
        if isinstance(v, (int, float)):
            return repr(v)
        return "'" + str(v).replace("'", "''") + "'"

    def insert(table, rows, cols):
        vals = ",\n  ".join("(" + ", ".join(["uid"] + [q(r[c]) for c in cols]) + ")" for r in rows)
        return f"insert into {table} (user_id, {', '.join(cols)}) values\n  {vals};\n"

    sql = ["-- Перенос данных из таблицы. Выполнить один раз после создания пользователя.",
           "do $$\ndeclare uid uuid := (select id from auth.users order by created_at limit 1);\nbegin",
           "if uid is null then raise exception 'Сначала создайте пользователя в Authentication → Users'; end if;",
           insert("printers", printers, ["id", "name", "price", "bought_on"]),
           insert("spools", spools, ["id", "name", "brand", "weight", "price", "bought_on"]),
           insert("extras", extras, ["id", "name", "category", "price", "bought_on"]),
           insert("orders", orders, ["id", "ordered_on", "plastic", "weight", "print_price", "model_price", "printer_id"]),
           "insert into settings (user_id, tariff) values (uid, 6.5) on conflict do nothing;",
           "end $$;"]
    (out / "seed.sql").write_text("\n".join(sql), encoding="utf-8")
    print(f"printers={len(printers)} spools={len(spools)} extras={len(extras)} orders={len(orders)} -> {out.resolve()}")


if __name__ == "__main__":
    main()
